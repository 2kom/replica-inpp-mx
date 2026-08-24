from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd
import pytest
from canasta_inpp.esquema import COLUMNAS_BASE, LAYOUTS_XLSX
from canasta_inpp.extraccion_xlsx import extraer_ponderadores
from canasta_inpp.utilidades import (
    guardar_csv,
    normalizar_columnas_con_codigo,
    normalizar_columnas_texto,
    normalizar_texto,
    normalizar_texto_con_codigo,
    resolver_sector_agrupado,
)

# -- guardar_csv: helper -------------------------------------------------


def _leer(ruta: Path) -> pd.DataFrame:
    return pd.read_csv(ruta, dtype=str, keep_default_na=False)


# -- normalizar_texto -------------------------------------------------------


@pytest.mark.parametrize(
    ("texto", "esperado"),
    [
        ("Árbol", "arbol"),
        ("ÁÉÍÓÚÜ", "aeiouu"),
        ("señor", "señor"),
        ("NIÑO", "niño"),
        ("hola, mundo!", "hola mundo"),
        ("hola   mundo", "hola mundo"),
        ("  hola  ", "hola"),
        ("11 Agricultura, cría y explotación", "11 agricultura cria y explotacion"),
    ],
)
def test_normalizar_texto(texto: str, esperado: str) -> None:
    assert normalizar_texto(texto) == esperado


# -- normalizar_columnas_texto -----------------------------------------------


def test_normalizar_columnas_texto_normaliza_solo_columnas_dadas() -> None:
    df = pd.DataFrame(
        {
            "generico": ["Soya y Otras Oleaginosas"],
            "sector": ["11 Agricultura, cría y explotación..."],
            "codigo": ["001"],
        }
    )
    resultado = normalizar_columnas_texto(df, ["generico", "sector"])
    assert resultado.loc[0, "generico"] == "soya y otras oleaginosas"
    assert resultado.loc[0, "sector"] == "11 agricultura cria y explotacion"
    assert resultado.loc[0, "codigo"] == "001"  # intacta, no pedida


def test_normalizar_columnas_texto_no_muta_el_df_original() -> None:
    df = pd.DataFrame({"generico": ["Frijol"]})
    normalizar_columnas_texto(df, ["generico"])
    assert df.loc[0, "generico"] == "Frijol"


def test_normalizar_columnas_texto_multiples_filas() -> None:
    df = pd.DataFrame({"generico": ["Soya", "Frijol", ""]})
    resultado = normalizar_columnas_texto(df, ["generico"])
    assert list(resultado["generico"]) == ["soya", "frijol", ""]


# -- normalizar_texto_con_codigo ---------------------------------------------
# Regresión: normalizar_texto (texto completo, sin distinguir código de
# nombre) destruye rangos SCIAN con guion -- "31-33 Industrias
# manufactureras" colapsa a "3133 industrias manufactureras", forma
# indistinguible de una rama real de 4 dígitos. Ver negociación
# data/negociaciones/2026-08-22-utilidades-normalizacion-scian.md.


@pytest.mark.parametrize(
    ("texto", "esperado"),
    [
        # rangos SCIAN oficiales (agrupación de sectores, no texto libre) --
        # el guion es parte del código, se preserva intacto
        ("31-33 Industrias manufactureras", "31-33 industrias manufactureras"),
        (
            "48-49 Transportes, correos y almacenamiento",
            "48-49 transportes correos y almacenamiento",
        ),
        # código simple, mismo resultado que normalizar_texto (sin guion que proteger)
        ("11 Agricultura, cría y explotación", "11 agricultura cria y explotacion"),
        # guion DENTRO del nombre (no en el código) -- sigue la regla general,
        # se pierde igual que en normalizar_texto (no es parte del código)
        ("31-33 Industria químico-farmacéutica", "31-33 industria quimicofarmaceutica"),
        # solo código, sin nombre -- se devuelve tal cual
        ("11", "11"),
        ("", ""),
    ],
)
def test_normalizar_texto_con_codigo(texto: str, esperado: str) -> None:
    assert normalizar_texto_con_codigo(texto) == esperado


# -- normalizar_columnas_con_codigo ------------------------------------------


def test_normalizar_columnas_con_codigo_preserva_rango_scian() -> None:
    df = pd.DataFrame(
        {
            "generico": ["Hilados de fibras artificiales y/o sintéticas"],
            "sector": ["31-33 Industrias manufactureras"],
            "codigo": ["123"],
        }
    )
    resultado = normalizar_columnas_con_codigo(df, ["sector"])
    assert resultado.loc[0, "sector"] == "31-33 industrias manufactureras"
    assert (
        resultado.loc[0, "generico"] == "Hilados de fibras artificiales y/o sintéticas"
    )  # intacta, no pedida
    assert resultado.loc[0, "codigo"] == "123"  # intacta, no pedida


def test_normalizar_columnas_con_codigo_no_muta_el_df_original() -> None:
    df = pd.DataFrame({"sector": ["31-33 Industrias manufactureras"]})
    normalizar_columnas_con_codigo(df, ["sector"])
    assert df.loc[0, "sector"] == "31-33 Industrias manufactureras"


def test_normalizar_columnas_con_codigo_multiples_columnas_y_filas() -> None:
    df = pd.DataFrame(
        {
            "sector": ["31-33 Industrias manufactureras", "11 Agricultura"],
            "subsector": ["48-49 Transportes", "111 Agricultura"],
        }
    )
    resultado = normalizar_columnas_con_codigo(df, ["sector", "subsector"])
    assert list(resultado["sector"]) == ["31-33 industrias manufactureras", "11 agricultura"]
    assert list(resultado["subsector"]) == ["48-49 transportes", "111 agricultura"]


# -- resolver_sector_agrupado -------------------------------------------------
# Regresión: sector "31-33"/"48-49" es un rango SCIAN (agrupa varios códigos de
# sector bajo un mismo nombre) -- dominio/calculo necesita el código concreto
# (31/32/33) para agrupar, no el rango. subsector siempre trae el código
# concreto en sus primeros 2 dígitos.


def test_resolver_sector_agrupado_con_nombre_combinado() -> None:
    df = pd.DataFrame(
        {
            "sector": ["31-33 Industrias manufactureras"],
            "subsector": ["311 Industria alimentaria"],
        }
    )
    resultado = resolver_sector_agrupado(df)
    assert resultado.loc[0, "sector"] == "31 Industrias manufactureras"


def test_resolver_sector_agrupado_bare_sin_nombre() -> None:
    # solo --ponderadores: sector/subsector son código bare, sin nombre pegado
    df = pd.DataFrame({"sector": ["48-49"], "subsector": ["492"]})
    resultado = resolver_sector_agrupado(df)
    assert resultado.loc[0, "sector"] == "49"


def test_resolver_sector_agrupado_no_toca_sector_sin_rango() -> None:
    df = pd.DataFrame(
        {
            "sector": ["11 Agricultura, cría y explotación"],
            "subsector": ["111 Agricultura"],
        }
    )
    resultado = resolver_sector_agrupado(df)
    assert resultado.loc[0, "sector"] == "11 Agricultura, cría y explotación"


def test_resolver_sector_agrupado_multiples_filas_mismo_rango_distinto_subsector() -> None:
    df = pd.DataFrame(
        {
            "sector": ["31-33 Industrias manufactureras", "31-33 Industrias manufactureras"],
            "subsector": ["311 Industria alimentaria", "331 Industrias metálicas básicas"],
        }
    )
    resultado = resolver_sector_agrupado(df)
    assert list(resultado["sector"]) == [
        "31 Industrias manufactureras",
        "33 Industrias manufactureras",
    ]


def test_resolver_sector_agrupado_no_muta_el_df_original() -> None:
    df = pd.DataFrame(
        {"sector": ["31-33 Industrias manufactureras"], "subsector": ["311 Industria alimentaria"]}
    )
    resolver_sector_agrupado(df)
    assert df.loc[0, "sector"] == "31-33 Industrias manufactureras"


# -- resolver_sector_agrupado: validación de jerarquía inconsistente ---------
# Regresión: sin validar, una jerarquía inconsistente producía una
# clasificación falsa en silencio en vez de fallar. Ver negociación
# data/negociaciones/2026-08-22-utilidades-normalizacion-scian.md § Seguimiento 2.


def test_resolver_sector_agrupado_lanza_si_subsector_fuera_del_rango() -> None:
    # subsector 481 (sector real 48) no pertenece al rango 31-33 declarado por sector
    df = pd.DataFrame(
        {
            "sector": ["31-33 Industrias manufactureras"],
            "subsector": ["481 Transporte aéreo"],
            "codigo": ["099"],
        }
    )
    with pytest.raises(ValueError, match="no pertenece al rango"):
        resolver_sector_agrupado(df)


def test_resolver_sector_agrupado_lanza_si_subsector_vacio() -> None:
    df = pd.DataFrame({"sector": ["48-49 Transportes"], "subsector": [""], "codigo": ["099"]})
    with pytest.raises(ValueError, match="no trae un código de 3 dígitos"):
        resolver_sector_agrupado(df)


def test_resolver_sector_agrupado_lanza_si_subsector_no_numerico() -> None:
    df = pd.DataFrame(
        {"sector": ["31-33 Industrias manufactureras"], "subsector": ["xx dato"], "codigo": ["099"]}
    )
    with pytest.raises(ValueError, match="no trae un código de 3 dígitos"):
        resolver_sector_agrupado(df)


def test_resolver_sector_agrupado_lanza_si_subsector_corto() -> None:
    # ortogonal a "vacío"/"no numérico": numérico y no vacío, pero longitud != 3 --
    # sin el chequeo de longitud, "31".isdigit() es True y el bug pasaría desapercibido
    df = pd.DataFrame(
        {"sector": ["31-33 Industrias manufactureras"], "subsector": ["31"], "codigo": ["099"]}
    )
    with pytest.raises(ValueError, match="no trae un código de 3 dígitos"):
        resolver_sector_agrupado(df)


def test_resolver_sector_agrupado_lanza_si_subsector_3_caracteres_no_numerico() -> None:
    # ortogonal a "vacío"/"no numérico": longitud exacta de 3, pero no numérico --
    # sin el chequeo de isdigit(), pasaría el largo y colaría un código no numérico
    df = pd.DataFrame(
        {
            "sector": ["31-33 Industrias manufactureras"],
            "subsector": ["31x dato"],
            "codigo": ["099"],
        }
    )
    with pytest.raises(ValueError, match="no trae un código de 3 dígitos"):
        resolver_sector_agrupado(df)


def test_resolver_sector_agrupado_mensaje_de_error_identifica_codigo_de_fila() -> None:
    df = pd.DataFrame(
        {
            "sector": ["31-33 Industrias manufactureras"],
            "subsector": ["481 Transporte aéreo"],
            "codigo": ["099"],
        }
    )
    with pytest.raises(ValueError, match="código 099"):
        resolver_sector_agrupado(df)


def test_resolver_sector_agrupado_mensaje_de_error_usa_indice_de_fila_sin_columna_codigo() -> None:
    # sin columna "codigo" en el df (ej. solo --ponderadores sin cruce) -- cae a índice de fila
    df = pd.DataFrame({"sector": ["31-33 Industrias manufactureras"], "subsector": ["481"]})
    with pytest.raises(ValueError, match="fila 0"):
        resolver_sector_agrupado(df)


# -- guardar_csv ---------------------------------------------------------


def test_guardar_csv_reindexa_a_columnas_base_y_rellena_columna_ausente(tmp_path: Path) -> None:
    df = pd.DataFrame({"generico": ["Soya"], "codigo": ["001"]})
    ruta = tmp_path / "salida.csv"
    guardar_csv(df, ruta, 2019)
    leido = _leer(ruta)
    assert list(leido.columns) == list(COLUMNAS_BASE)
    assert leido.loc[0, "generico"] == "Soya"
    assert leido.loc[0, "produccion total"] == ""  # columna entera ausente -> vacio, no "-"


def test_guardar_csv_preserva_cero_real_de_ponderador(tmp_path: Path) -> None:
    # 0 es un valor real (el generico no participa en ese destino/etapa), no una ausencia
    df = pd.DataFrame({"generico": ["Arroz"], "exportaciones": ["0"]})
    ruta = tmp_path / "salida.csv"
    guardar_csv(df, ruta, 2019)
    leido = _leer(ruta)
    assert leido.loc[0, "exportaciones"] == "0"


def test_guardar_csv_preserva_string_exacto_de_ponderador(tmp_path: Path) -> None:
    # precision cruda del xlsx (notacion cientifica) no debe convertirse a float
    df = pd.DataFrame({"generico": ["Soya"], "produccion total": ["3.0944225043218539E-2"]})
    ruta = tmp_path / "salida.csv"
    guardar_csv(df, ruta, 2019)
    leido = _leer(ruta)
    assert leido.loc[0, "produccion total"] == "3.0944225043218539E-2"


def test_guardar_csv_celda_nan_en_columna_presente_se_guarda_como_guion(tmp_path: Path) -> None:
    # "N/A" de INEGI en encadenamiento_exportacion, ya convertido a NaN por
    # extraer_encadenamiento -- celda puntual sin dato, no columna ausente
    df = pd.DataFrame(
        {
            "generico": ["Soya", "Frijol"],
            "encadenamiento exportacion": ["1.416519", float("nan")],
        }
    )
    ruta = tmp_path / "salida.csv"
    guardar_csv(df, ruta, 2025)
    leido = _leer(ruta)
    assert leido.loc[0, "encadenamiento exportacion"] == "1.416519"
    assert leido.loc[1, "encadenamiento exportacion"] == "-"


def test_guardar_csv_lanza_valueerror_si_ponderador_tiene_nan() -> None:
    # regresion: fillna("-") universal confundia un ponderador vacio (dato
    # requerido, nunca deberia faltar) con el N/A legitimo de encadenamiento.
    # ver negociacion data/negociaciones/2026-08-22-utilidades-normalizacion-scian.md
    # § "Seguimiento 5"
    df = pd.DataFrame(
        {"generico": ["Soya", "Frijol"], "codigo": ["001", "002"], "exportaciones": ["1.5", None]}
    )
    with pytest.raises(ValueError, match="'exportaciones' trae 1 celda"):
        guardar_csv(df, Path("/tmp/no-deberia-escribirse.csv"), 2019)


def test_guardar_csv_lanza_valueerror_si_generico_tiene_nan() -> None:
    # no solo ponderadores -- cualquier columna fuera de encadenamiento
    df = pd.DataFrame({"generico": ["Soya", None], "codigo": ["001", "002"]})
    with pytest.raises(ValueError, match="'generico'"):
        guardar_csv(df, Path("/tmp/no-deberia-escribirse.csv"), 2019)


def test_guardar_csv_mensaje_de_error_identifica_codigos_afectados() -> None:
    df = pd.DataFrame({"generico": ["Soya", "Frijol", "Garbanzo"], "codigo": ["001", "002", "003"]})
    df.loc[[1, 2], "generico"] = None
    with pytest.raises(ValueError, match=r"\['002', '003'\]"):
        guardar_csv(df, Path("/tmp/no-deberia-escribirse.csv"), 2019)


def test_guardar_csv_no_lanza_si_solo_columnas_de_encadenamiento_tienen_nan(
    tmp_path: Path,
) -> None:
    # sanity check inverso: NaN permitido en encadenamiento no debe disparar el ValueError
    df = pd.DataFrame(
        {
            "generico": ["Soya"],
            "encadenamiento exportacion": [float("nan")],
            "encadenamiento uso final": [float("nan")],
        }
    )
    ruta = tmp_path / "salida.csv"
    guardar_csv(df, ruta, 2025)  # no debe lanzar
    leido = _leer(ruta)
    assert leido.loc[0, "encadenamiento exportacion"] == "-"
    assert leido.loc[0, "encadenamiento uso final"] == "-"


# -- guardar_csv: encadenamiento total/produccion nacional NUNCA admiten N/A -
# Regresión: COLUMNAS_ENCADENAMIENTO (las 4) se usaba como excepción completa
# -- total/produccion_nacional NUNCA traen N/A en el xlsx real (confirmado en
# extraer_encadenamiento), así que un NaN ahí es defecto de extracción, mismo
# tratamiento que ponderadores. Ver negociación
# data/negociaciones/2026-08-22-utilidades-normalizacion-scian.md § "Seguimiento 6".


@pytest.mark.parametrize("columna", ["encadenamiento total", "encadenamiento produccion nacional"])
def test_guardar_csv_lanza_valueerror_si_encadenamiento_total_o_produccion_nacional_tiene_nan(
    columna: str,
) -> None:
    df = pd.DataFrame({"generico": ["Soya"], "codigo": ["001"], columna: [float("nan")]})
    with pytest.raises(ValueError, match=f"'{columna}'"):
        guardar_csv(df, Path("/tmp/no-deberia-escribirse.csv"), 2025)


# -- guardar_csv: fallback a posición cuando no hay codigo utilizable -------
# Decisión de diseño: reportar POSICIÓN dentro del df (0-indexed) SIEMPRE, por
# contrato, nunca el índice de `df` -- sin importar qué índice traiga `df`.
# Los 3 extractores actuales (extraer_ponderadores/extraer_canasta/
# extraer_encadenamiento) siempre devuelven índice fresco, y `df.merge(...)`
# (el llamador real en `generar_canasta.py::main()`) también resetea a un
# RangeIndex fresco -- pero reportar por posición evita toda la clase de bugs
# de índice (etiquetas duplicadas rompiendo `.loc`, tipos numpy en el mensaje)
# sin depender de esa invariante. Ver negociación
# data/negociaciones/2026-08-22-utilidades-normalizacion-scian.md § "Seguimiento 10".


def test_guardar_csv_mensaje_de_error_usa_posicion_cuando_no_hay_columna_codigo() -> None:
    # sin "codigo" en el df original (ej. solo --ponderadores sin cruce) --
    # el fallback usa la posición, no el "codigo" vacío que deja el reindex
    df = pd.DataFrame({"generico": ["Soya", None]})
    with pytest.raises(ValueError, match=r"\[1\]"):
        guardar_csv(df, Path("/tmp/no-deberia-escribirse.csv"), 2019)


def test_guardar_csv_mensaje_de_error_usa_posicion_si_codigo_es_none_en_la_misma_fila() -> None:
    # codigo SÍ existe como columna, pero justo la fila con el problema tiene codigo=None --
    # no sirve como identificador, cae a la posición igual que si no hubiera columna codigo
    df = pd.DataFrame({"generico": ["Soya", None], "codigo": ["001", None]})
    with pytest.raises(ValueError, match=r"\[1\]"):
        guardar_csv(df, Path("/tmp/no-deberia-escribirse.csv"), 2019)


def test_guardar_csv_mensaje_de_error_usa_posicion_si_codigo_es_vacio_en_la_misma_fila() -> None:
    df = pd.DataFrame({"generico": ["Soya", None], "codigo": ["001", ""]})
    with pytest.raises(ValueError, match=r"\[1\]"):
        guardar_csv(df, Path("/tmp/no-deberia-escribirse.csv"), 2019)


def test_guardar_csv_mensaje_de_error_usa_codigo_cuando_esta_disponible_en_la_fila() -> None:
    # sanity check inverso: si codigo SÍ tiene valor en la fila afectada, se usa (no la posición)
    df = pd.DataFrame({"generico": ["Soya", "Frijol", None], "codigo": ["001", "002", "003"]})
    with pytest.raises(ValueError, match=r"\['003'\]"):
        guardar_csv(df, Path("/tmp/no-deberia-escribirse.csv"), 2019)


def test_guardar_csv_mensaje_de_error_con_indice_duplicado_no_rompe() -> None:
    # regresion: .loc[etiqueta] con índice duplicado devuelve una Series, no
    # un escalar -- rompe pd.notna(...) con "the truth value of a Series is
    # ambiguous" en vez del ValueError contractual. Enumerar por posición
    # (nunca .loc por etiqueta) es inmune, sin importar duplicados.
    df = pd.DataFrame(
        {"generico": ["Soya", None], "codigo": ["001", "002"]},
        index=[7, 7],
    )
    with pytest.raises(ValueError, match=r"\['002'\]"):
        guardar_csv(df, Path("/tmp/no-deberia-escribirse.csv"), 2019)


def test_guardar_csv_mensaje_de_error_ignora_indice_custom_reporta_posicion() -> None:
    # con índice no trivial [10, 20], se reporta la posición (1), no la
    # etiqueta (20) -- ver nota de diseño arriba del bloque.
    df = pd.DataFrame({"generico": ["Soya", None]}, index=[10, 20])
    with pytest.raises(ValueError, match=r"\[1\]"):
        guardar_csv(df, Path("/tmp/no-deberia-escribirse.csv"), 2019)


def test_guardar_csv_integrado_extraer_ponderadores_con_peso_vacio_lanza(tmp_path: Path) -> None:
    # prueba integrada pedida por el evaluador: extraer_ponderadores -> guardar_csv,
    # con un xlsx real donde una celda de peso viene vacia (None) -- la igualdad
    # de codigos entre hojas NO garantiza que cada celda tenga valor
    layout = LAYOUTS_XLSX[2019]

    def _fila(codigo: str, nombre: str, peso: object) -> tuple:
        return (None, 11, 111, 1111, 11111, 111110, codigo, nombre, peso)

    wb = openpyxl.Workbook()
    hoja_por_defecto = wb.active
    assert hoja_por_defecto is not None
    wb.remove(hoja_por_defecto)

    for hoja in (
        layout.hoja_produccion_total,
        layout.hoja_bienes_intermedios,
        layout.hoja_bienes_finales,
        layout.hoja_exportaciones,
    ):
        ws = wb.create_sheet(hoja)
        ws.append(_fila("001", "Soya y otras oleaginosas", 10.0))

    # demanda interna trae 3 pesos por fila -- el segundo (consumo) queda vacio
    ws = wb.create_sheet(layout.hoja_demanda_interna)
    ws.append(
        (None, 11, 111, 1111, 11111, 111110, "001", "Soya y otras oleaginosas", 40.0, None, 42.0)
    )

    ruta_xlsx = tmp_path / "ponderadores_2019.xlsx"
    wb.save(ruta_xlsx)

    df = extraer_ponderadores(ruta_xlsx, 2019)
    assert df.loc[0, "demanda interna consumo"] is None  # confirma el hueco antes de guardar_csv

    with pytest.raises(ValueError, match="'demanda interna consumo'"):
        guardar_csv(df, tmp_path / "salida.csv", 2019)


def test_guardar_csv_distingue_columna_ausente_de_celda_nan_en_la_misma_corrida(
    tmp_path: Path,
) -> None:
    # ambos casos a la vez: encadenamiento_exportacion presente con un NaN puntual
    # ("-"), encadenamiento_uso_final ausente por completo ("")
    df = pd.DataFrame(
        {
            "generico": ["Soya"],
            "encadenamiento exportacion": [float("nan")],
        }
    )
    ruta = tmp_path / "salida.csv"
    guardar_csv(df, ruta, 2025)
    leido = _leer(ruta)
    assert leido.loc[0, "encadenamiento exportacion"] == "-"
    assert leido.loc[0, "encadenamiento uso final"] == ""


def test_guardar_csv_columna_extra_no_declarada_se_descarta(tmp_path: Path) -> None:
    df = pd.DataFrame({"generico": ["Soya"], "columna_inventada": ["x"]})
    ruta = tmp_path / "salida.csv"
    guardar_csv(df, ruta, 2019)
    leido = _leer(ruta)
    assert "columna_inventada" not in leido.columns


def test_guardar_csv_columna_extra_advierte_sin_lanzar(
    capsys: pytest.CaptureFixture[str], tmp_path: Path
) -> None:
    df = pd.DataFrame({"generico": ["Soya"], "ponderdor": ["10.5"]})
    ruta = tmp_path / "salida.csv"
    guardar_csv(df, ruta, 2019)  # no debe lanzar
    salida = capsys.readouterr().out
    assert "ponderdor" in salida


def test_guardar_csv_preserva_orden_de_filas_y_columnas_multirregistro(tmp_path: Path) -> None:
    df = pd.DataFrame(
        {
            "generico": ["soya", "frijol", "tortilla"],
            "codigo": ["001", "002", "003"],
            "sector": ["11 agricultura", "11 agricultura", "31 industrias"],
        }
    )
    ruta = tmp_path / "salida.csv"
    guardar_csv(df, ruta, 2019)
    leido = _leer(ruta)
    assert list(leido["generico"]) == ["soya", "frijol", "tortilla"]
    assert list(leido["codigo"]) == ["001", "002", "003"]
    assert (leido["produccion total"] == "").all()  # columna ausente en las 3 filas
